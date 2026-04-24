from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from app.domain.magento_columns import STAFF_TEMPLATE_COLUMNS


class ExcelTemplateGenerator:
    def generate(self, output_path: Path, dropdowns: dict[str, list[str]] | None = None) -> Path:
        dropdowns = dropdowns or {}
        output_path.parent.mkdir(parents=True, exist_ok=True)

        wb = Workbook()
        ws = wb.active
        ws.title = "Product_Upload"

        required_fill = PatternFill("solid", fgColor="FFE5E5")
        header_font = Font(bold=True)

        for col_idx, column in enumerate(STAFF_TEMPLATE_COLUMNS, start=1):
            cell = ws.cell(row=1, column=col_idx, value=column)
            cell.font = header_font
            if column in {
                "sku",
                "name",
                "gemstone",
                "origin",
                "treatment",
                "carat_weight",
                "shape",
                "colour",
                "cut",
                "dimensions",
                "certification",
                "price",
                "qty",
                "hsn_code",
                "dispatch_days",
                "shipping_days",
                "return_policy",
            }:
                cell.fill = required_fill

        ws.append(
            [
                "GP159106",
                "GP159106",
                "Amethyst - 5.11 Carats",
                "Amethyst",
                "Brazil",
                "Unheated and Untreated (No Indications Observed)",
                5.11,
                "Oval",
                "Violet",
                "Faceted",
                "Faceted",
                "13.87x9.68x6.68 mm",
                "Not Calibrated",
                "AGR Certified",
                "",
                5500,
                "",
                1,
                "GP",
                "133045",
                "71039949",
                "0 Business Days (+5 days for typical Jewellery; +14 days for Bracelet/Gold/CZ/Diamond Items)",
                0,
                "10 Day Money-Back Returns*",
                "/g/p/gp159106-1-220426.jpg",
                "",
                "",
                "",
                "gp159106-220426.mp4",
                "",
                "",
                "",
                "",
                "",
                "Amethyst",
                "MP",
                "provides balance and stability",
            ]
        )

        values_ws = wb.create_sheet("Dropdown_Values")
        for col_idx, (field, values) in enumerate(dropdowns.items(), start=1):
            values_ws.cell(row=1, column=col_idx, value=field).font = header_font
            for row_idx, value in enumerate(values, start=2):
                values_ws.cell(row=row_idx, column=col_idx, value=value)

            if values:
                col_letter = values_ws.cell(row=1, column=col_idx).column_letter
                formula = f"=Dropdown_Values!${col_letter}$2:${col_letter}${len(values)+1}"
                validation = DataValidation(type="list", formula1=formula, allow_blank=True)
                ws.add_data_validation(validation)
                if field in STAFF_TEMPLATE_COLUMNS:
                    target_col = STAFF_TEMPLATE_COLUMNS.index(field) + 1
                    target_letter = ws.cell(row=1, column=target_col).column_letter
                    validation.add(f"{target_letter}2:{target_letter}1000")

        instructions = wb.create_sheet("Instructions")
        instructions["A1"] = "Rules"
        instructions["A1"].font = header_font
        instructions["A2"] = "Red headers are required for loose gemstone products."
        instructions["A3"] = "Use SKU-based image naming wherever possible."
        instructions["A4"] = "Do not manually fill Magento-only fields such as status, tax_class_id, small_image, thumbnail."
        instructions["A5"] = "Upload products as draft/disabled first, then review before publishing."

        wb.save(output_path)
        return output_path
